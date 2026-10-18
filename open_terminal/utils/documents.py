"""Document text extraction utilities.

Extracts readable text from binary document formats so LLMs can
consume their content.  Each ``extract_*`` function takes a file path
and returns the document's text as a plain string.

All libraries used are permissively licensed (MIT / BSD).
"""

import zipfile
import os
from functools import wraps

from open_terminal.utils.limits import (
    MAX_DOCUMENT_BYTES, MAX_DOCUMENT_CHARS, MAX_DOCUMENT_CELLS,
    MAX_TEXT_LINES, ResourceLimitError,
)


class TextParts(list):
    def __init__(self):
        super().__init__()
        self.characters = 0

    def append(self, value):
        self.characters += len(value) + 2
        if self.characters > MAX_DOCUMENT_CHARS or len(self) >= MAX_TEXT_LINES:
            raise ResourceLimitError("Document text exceeds the extraction limit.")
        super().append(value)


def bounded_extractor(function):
    @wraps(function)
    def extract(path, *args, **kwargs):
        if os.path.getsize(path) > MAX_DOCUMENT_BYTES:
            raise ResourceLimitError("Document exceeds the input size limit.")
        if zipfile.is_zipfile(path):
            with zipfile.ZipFile(path) as archive:
                if len(archive.infolist()) > 10_000 or sum(info.file_size for info in archive.infolist()) > MAX_DOCUMENT_BYTES:
                    raise ResourceLimitError("Expanded document exceeds the extraction limit.")
        text = function(path, *args, **kwargs)
        if len(text) > MAX_DOCUMENT_CHARS:
            raise ResourceLimitError("Document text exceeds the extraction limit.")
        return text
    return extract



@bounded_extractor
def extract_pdf(file_path: str, strict: bool = False) -> str:
    """Extract text from a PDF file."""
    from pypdf import PdfReader

    reader = PdfReader(file_path)
    if strict and reader.is_encrypted:
        raise ValueError("Encrypted PDF. Provide an unencrypted copy.")
    pages = TextParts()
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    missing = [str(i) for i, text in enumerate(pages, 1) if not text.strip()]
    if strict and missing:
        raise ValueError(
            f"No extractable text on PDF pages {', '.join(missing)}. OCR is required before comparison."
        )
    return "\n".join(pages)


@bounded_extractor
def extract_docx(file_path: str) -> str:
    """Extract current paragraph/table text, including tracked insertions, without saving."""
    from docx import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    from docx.oxml.ns import qn
    from lxml import etree

    doc = DocxDocument(file_path)
    body = doc.element.body
    # Ignore historical properties before finding revisions in the current text.
    for node in list(body.iter()):
        if isinstance(node.tag, str) and node.tag.startswith(qn("w:")) and node.tag.endswith("Change"):
            node.getparent().remove(node)
    for node in body.xpath(".//w:tr[w:trPr/w:del] | .//w:tc[w:tcPr/w:cellDel]"):
        node.getparent().remove(node)
    joined_paragraphs = body.xpath(".//w:p[w:pPr/w:rPr/w:del]")
    etree.strip_elements(body, qn("w:del"), qn("w:moveFrom"), with_tail=False)
    etree.strip_tags(body, qn("w:ins"), qn("w:moveTo"))
    # A deleted paragraph mark joins its text to the following paragraph.
    for paragraph in reversed(joined_paragraphs):
        following = paragraph.getnext()
        if following is not None and following.tag == qn("w:p"):
            for child in list(following):
                if child.tag != qn("w:pPr"):
                    paragraph.append(child)
            following.getparent().remove(following)
    parts = TextParts()
    for child in doc.element.body:
        if child.tag.endswith("}p"):
            parts.append(Paragraph(child, doc).text)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, doc).rows:
                parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


@bounded_extractor
def extract_xlsx(file_path: str) -> str:
    """Extract text from an Excel (.xlsx) file."""
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    parts = TextParts()
    cells = 0
    try:
        for sheet in wb.worksheets:
            cells += (sheet.max_row or 0) * (sheet.max_column or 0)
            if cells > MAX_DOCUMENT_CELLS:
                raise ResourceLimitError("Spreadsheet cell range exceeds the extraction limit.")
            parts.append(f"--- {sheet.title} ---")
            visited = 0
            for row in sheet.iter_rows(values_only=True):
                visited += len(row)
                if visited > MAX_DOCUMENT_CELLS:
                    raise ResourceLimitError("Spreadsheet cell range exceeds the extraction limit.")
                parts.append("\t".join(str(c) if c is not None else "" for c in row))
        return "\n".join(parts)
    finally:
        wb.close()


@bounded_extractor
def extract_pptx(file_path: str) -> str:
    """Extract text from a PowerPoint (.pptx) file."""
    from pptx import Presentation

    prs = Presentation(file_path)
    parts = TextParts()
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


@bounded_extractor
def extract_rtf(file_path: str) -> str:
    """Extract text from a Rich Text Format (.rtf) file."""
    from striprtf.striprtf import rtf_to_text

    with open(file_path, "rb") as f:
        raw = f.read()
    return rtf_to_text(raw.decode("utf-8", errors="replace"))


@bounded_extractor
def extract_xls(file_path: str) -> str:
    """Extract text from a legacy Excel (.xls) file."""
    import xlrd

    wb = xlrd.open_workbook(file_path)
    parts = TextParts()
    if sum(sheet.nrows * sheet.ncols for sheet in wb.sheets()) > MAX_DOCUMENT_CELLS:
        raise ResourceLimitError("Spreadsheet cell range exceeds the extraction limit.")
    for sheet in wb.sheets():
        parts.append(f"--- {sheet.name} ---")
        for row_idx in range(sheet.nrows):
            parts.append("\t".join(
                str(sheet.cell_value(row_idx, col_idx))
                for col_idx in range(sheet.ncols)
            ))
    return "\n".join(parts)


@bounded_extractor
def extract_odt(file_path: str) -> str:
    """Extract text from an OpenDocument Text (.odt) file."""
    from lxml import etree

    with zipfile.ZipFile(file_path) as zf:
        with zf.open("content.xml") as f:
            tree = etree.parse(f)
    ns = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
    return "\n".join(
        "".join(p.itertext())
        for p in tree.iter(f"{{{ns}}}p")
    )


@bounded_extractor
def extract_ods(file_path: str) -> str:
    """Extract text from an OpenDocument Spreadsheet (.ods) file."""
    from lxml import etree

    with zipfile.ZipFile(file_path) as zf:
        with zf.open("content.xml") as f:
            tree = etree.parse(f)
    ns_table = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"
    ns_text = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
    parts = TextParts()
    for table in tree.iter(f"{{{ns_table}}}table"):
        name = table.get(f"{{{ns_table}}}name", "Sheet")
        parts.append(f"--- {name} ---")
        for row in table.iter(f"{{{ns_table}}}table-row"):
            cells = []
            for cell in row.iter(f"{{{ns_table}}}table-cell"):
                cell_text = " ".join(
                    "".join(p.itertext())
                    for p in cell.iter(f"{{{ns_text}}}p")
                )
                cells.append(cell_text)
            parts.append("\t".join(cells))
    return "\n".join(parts)


@bounded_extractor
def extract_odp(file_path: str) -> str:
    """Extract text from an OpenDocument Presentation (.odp) file."""
    from lxml import etree

    with zipfile.ZipFile(file_path) as zf:
        with zf.open("content.xml") as f:
            tree = etree.parse(f)
    ns_draw = "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
    ns_text = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
    parts = TextParts()
    for i, page in enumerate(tree.iter(f"{{{ns_draw}}}page"), 1):
        parts.append(f"--- Slide {i} ---")
        for p in page.iter(f"{{{ns_text}}}p"):
            text = "".join(p.itertext()).strip()
            if text:
                parts.append(text)
    return "\n".join(parts)


@bounded_extractor
def extract_epub(file_path: str, strict: bool = False) -> str:
    """Extract text from an EPUB e-book."""
    from lxml import etree

    parts = TextParts()
    with zipfile.ZipFile(file_path) as zf:
        # Parse the container to find the root file
        with zf.open("META-INF/container.xml") as cf:
            container = etree.parse(cf)
        ns_container = "urn:oasis:names:tc:opendocument:xmlns:container"
        rootfile = container.find(f".//{{{ns_container}}}rootfile")
        if rootfile is None:
            rootfile = container.xpath("//*[local-name()='rootfile']")
            rootfile = rootfile[0] if rootfile else None

        if rootfile is not None:
            opf_path = rootfile.get("full-path", "")
            opf_dir = opf_path.rsplit("/", 1)[0] + "/" if "/" in opf_path else ""
            with zf.open(opf_path) as opf_file:
                opf = etree.parse(opf_file)
            spine_ids = [
                item.get("idref")
                for item in opf.xpath("//*[local-name()='itemref']")
            ]
            manifest = {
                item.get("id"): item.get("href")
                for item in opf.xpath("//*[local-name()='item']")
            }
            for idref in spine_ids:
                href = manifest.get(idref, "")
                item_path = opf_dir + href if not href.startswith("/") else href.lstrip("/")
                try:
                    with zf.open(item_path) as html_file:
                        html_tree = etree.parse(html_file, etree.HTMLParser())
                        body = html_tree.find(".//body")
                        if body is not None:
                            text = "".join(body.itertext())
                            parts.append(text.strip())
                except (KeyError, etree.XMLSyntaxError):
                    if strict:
                        raise ValueError(f"Could not extract EPUB chapter: {item_path}")
                    continue
        else:
            for name in zf.namelist():
                if name.endswith((".html", ".xhtml", ".htm")):
                    try:
                        with zf.open(name) as html_file:
                            html_tree = etree.parse(html_file, etree.HTMLParser())
                            body = html_tree.find(".//body")
                            if body is not None:
                                text = "".join(body.itertext())
                                parts.append(text.strip())
                    except etree.XMLSyntaxError:
                        if strict:
                            raise ValueError(f"Could not extract EPUB chapter: {name}")
                        continue
    return "\n\n".join(parts)


@bounded_extractor
def extract_eml(file_path: str) -> str:
    """Extract text from an email message (.eml)."""
    import email
    from email import policy

    with open(file_path, "rb") as f:
        msg = email.message_from_binary_file(f, policy=policy.default)
    parts = TextParts()
    for header in ("From", "To", "Cc", "Date", "Subject"):
        val = msg.get(header)
        if val:
            parts.append(f"{header}: {val}")
    parts.append("")  # blank line after headers
    body = msg.get_body(preferencelist=("plain", "html"))
    if body:
        content = body.get_content()
        if body.get_content_type() == "text/html":
            from lxml import etree
            tree = etree.HTML(content)
            content = "".join(tree.itertext()) if tree is not None else content
        parts.append(content)
    return "\n".join(parts)


# MIME type / extension → extractor mapping.
# Checked in order by read_file; the first match wins.
# Each entry: (mime_type_or_None, file_extension_or_None, extractor)
EXTRACTORS: list[tuple[str | None, str | None, callable]] = [
    ("application/pdf", None, extract_pdf),
    ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", None, extract_docx),
    ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", None, extract_xlsx),
    ("application/vnd.openxmlformats-officedocument.presentationml.presentation", None, extract_pptx),
    ("application/rtf", ".rtf", extract_rtf),
    ("application/vnd.ms-excel", ".xls", extract_xls),
    ("application/vnd.oasis.opendocument.text", ".odt", extract_odt),
    ("application/vnd.oasis.opendocument.spreadsheet", ".ods", extract_ods),
    ("application/vnd.oasis.opendocument.presentation", ".odp", extract_odp),
    ("application/epub+zip", ".epub", extract_epub),
    ("message/rfc822", ".eml", extract_eml),
]


if __name__ == "__main__":
    import json
    import mimetypes
    import sys

    try:
        path = json.load(sys.stdin)["path"]
        mime = mimetypes.guess_type(path)[0]
        extractor = next(fn for kind, suffix, fn in EXTRACTORS if kind and kind == mime or suffix and path.lower().endswith(suffix))
        result = {"text": extractor(path)}
    except Exception as exc:
        result = {"error": str(exc)[:1000]}
    json.dump(result, sys.stdout)
